import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

d = Document(sys.argv[1])
st = d.styles
NAVY = RGBColor(0x1F, 0x3A, 0x5F); TEAL = RGBColor(0x1B, 0x6E, 0x7A); GREY = RGBColor(0x44, 0x44, 0x44)

def S(name):
    for s in st:
        if s.name == name: return s
    raise KeyError(name)

def font(style, name, size=None, bold=None, color=None, italic=None):
    s = S(style); f = s.font
    f.name = name
    rpr = s.element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts'); rpr.append(rfonts)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rfonts.set(qn(a), name)
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rfonts.get(qn(a)) is not None:
            del rfonts.attrib[qn(a)]
    if size: f.size = Pt(size)
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color is not None: f.color.rgb = color

for name in ['Normal', 'Body Text', 'First Paragraph', 'Compact']:
    if name in [s.name for s in st]:
        font(name, 'Calibri', 10.5, color=RGBColor(0x22, 0x22, 0x22))
        pf = S(name).paragraph_format; pf.space_after = Pt(5); pf.line_spacing = 1.12

font('Title', 'Calibri', 28, True, NAVY)
font('Subtitle', 'Calibri', 15, False, TEAL)
for n in ['Author', 'Date']:
    font(n, 'Calibri', 11, False, GREY)
font('Heading 1', 'Calibri', 24, True, NAVY)
font('Heading 2', 'Calibri', 19, True, NAVY)
font('Heading 3', 'Calibri', 14, True, TEAL)
font('Heading 4', 'Calibri', 12, True, NAVY)
font('Heading 5', 'Calibri', 11, True, TEAL)
for h in ['Heading 1', 'Heading 2']:
    S(h).paragraph_format.page_break_before = True
    S(h).paragraph_format.space_after = Pt(10)
for h in ['Heading 3']:
    S(h).paragraph_format.space_before = Pt(16); S(h).paragraph_format.space_after = Pt(6)
    # bottom border under section headings
    ppr = S(h).element.get_or_add_pPr()
    bdr = OxmlElement('w:pBdr'); b = OxmlElement('w:bottom')
    for k, v in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '2'), ('w:color', '1B6E7A')):
        b.set(qn(k), v)
    bdr.append(b); ppr.append(bdr)
for h in ['Heading 4', 'Heading 5']:
    S(h).paragraph_format.space_before = Pt(10); S(h).paragraph_format.space_after = Pt(4)

# code blocks
font('Source Code', 'Consolas', 7.5)
sc = S('Source Code'); pf = sc.paragraph_format
pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.0
ppr = sc.element.get_or_add_pPr()
shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), 'F4F6F8')
ppr.append(shd)
# inline code
if 'Verbatim Char' in [s.name for s in st]:
    font('Verbatim Char', 'Consolas', 7.5)
# block quote
if 'Block Text' in [s.name for s in st]:
    font('Block Text', 'Calibri', 10.5, italic=False, color=NAVY)

for sec in d.sections:
    sec.page_width = Inches(8.5); sec.page_height = Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.7)
    sec.top_margin = sec.bottom_margin = Inches(0.8)

# tables: borders + shaded header row + smaller font
def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    tcPr.append(sh)
for t in d.tables:
    tblPr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}')
        e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:space'), '0'); e.set(qn('w:color'), 'B8C4CE')
        borders.append(e)
    tblPr.append(borders)
    for i, row in enumerate(t.rows):
        for cell in row.cells:
            if i == 0:
                set_cell_shading(cell, 'DCE6EE')
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(9)
                    if i == 0: r.font.bold = True
d.save(sys.argv[2])
